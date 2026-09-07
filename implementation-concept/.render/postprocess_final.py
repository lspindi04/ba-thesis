from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement

path = Path(r"C:\zenlab-Projects\ba-thesis\implementation-concept\DeviceIntegrationPlugin_Kurzkonzept_geprueft.docx")
doc = Document(path)

replacements = {
    "Analog zum bestehenden DeviceManagementPlugin werden zwei getrennte code-first-gRPC-Servicegrenzen verwendet: eine für Client-Aufrufe und eine für Agents beziehungsweise Integrationen. Der jeweilige SmartProxy ist bereits der konkrete gRPC-Client und implementiert denselben Servicevertrag wie der serverseitige Service.":
    "Analog zum bestehenden DeviceManagementPlugin werden zwei getrennte code-first-gRPC-Servicegrenzen verwendet: eine für Client-Aufrufe und eine für Agents beziehungsweise Integrationen. Der jeweilige SmartProxy ist der konkrete gRPC-Client und implementiert denselben Vertrag wie der serverseitige Service. Für langlebige Node- und Duplex-Streams konfigurieren beide Builder ausdrücklich WithOneShotClient. Der vorhandene Client<T>.GetStream ist wegen automatischem Reconnect und unbounded Zwischen-Channel dafür ungeeignet; OneShotClient.GetStream wird deshalb ohne Retry und mit try/finally für den gRPC-Channel abgesichert.",
    "Der DeviceIntegrationAgentRuntime registriert sich über den DeviceIntegrationAgentServiceProxy beim DeviceIntegrationAgentService und sendet RequestedServiceId, Metadaten sowie die vollständige aktuelle Device-Liste.":
    "Der DeviceIntegrationAgentRuntime registriert sich über den DeviceIntegrationAgentServiceProxy beim DeviceIntegrationAgentService und sendet RequestedServiceId, Metadaten sowie die vollständige aktuelle Device-Liste. Jeder Snapshot-Eintrag trägt einen im Agent-Prozess stabilen AgentDeviceKey.",
    "NotifyDeviceList enthält immer die vollständige aktuelle Device-Liste, ersetzt den Snapshot atomar und erneuert zugleich die Lease. Ein separates Agent-KeepAlive gibt es nicht.":
    "NotifyDeviceList enthält immer die vollständige aktuelle Device-Liste, ersetzt den Snapshot atomar und erneuert zugleich die Lease. Die Antwort liefert jedes Mal die vollständige Zuordnung AgentDeviceKey zu kanonischer DeviceId; der Agent ersetzt seine bidirektionale Zuordnung atomar. Ein separates Agent-KeepAlive gibt es nicht.",
    "DeviceIntegrationClientService und DeviceIntegrationAgentService laufen im Plugin. Sie greifen ausschließlich über Lifecycle-, Dispatcher-, Connection- und Subscription-Abstraktionen auf die gemeinsam als SingleInstance registrierten Zustände zu.":
    "DeviceIntegrationClientService und DeviceIntegrationAgentService laufen im Plugin. Sie greifen ausschließlich über Lifecycle-, Dispatcher-, Connection- und Subscription-Abstraktionen auf die gemeinsam als SingleInstance registrierten Zustände zu. Weil ein vorhandener DeviceAgent SubscriptionCreated bereits synchron vor seiner MethodResponse melden kann, puffert der SubscriptionManager diese Reihenfolge im Pending-Zustand begrenzt und gibt an Clients stets Accepted oder Rejected vor Created aus.",
    "Client- und Agent-Serviceverträge, DTOs, beide SmartProxy-Implementierungen und ClientBuilder.":
    "Client- und Agent-Serviceverträge, DTOs, SmartProxys, gemeinsamer MessageCodec sowie Builder mit expliziter OneShot-Konfiguration für langlebige Streams.",
    "AgentRuntime, SmartProxy, Outbox, WorkScheduler, Mapper sowie Snapshot- und Event-Bridge für bestehende DeviceAgents.":
    "AgentRuntime, SmartProxy und ein frischer AgentSessionRuntime je RegistrationId mit eigener Outbox, Cancellation, WorkScheduler, DeviceId-Zuordnung sowie Snapshot- und Event-Bridge.",
    "Bestehenden Router mit Characterization Tests absichern; Abhängigkeiten dokumentieren.":
    "Zuerst die begonnene Solution buildfähig und als Plugin auffindbar machen: GlobalAssemblyInfo-Pfad korrigieren sowie MEF-Export und PluginMetadata ergänzen. Danach den vorhandenen Router mit Characterization Tests absichern und seine Abhängigkeiten dokumentieren.",
    "Client-/Agent-Serviceverträge, beide SmartProxys, Register, Lease und Cleanup umsetzen.":
    "Client-/Agent-Serviceverträge, beide SmartProxys, OneShot-Streaming, atomare Aktivierungsbarriere, Register, Vollsnapshot-Lease und idempotentes Cleanup umsetzen.",
    "Service-/Device-/Pfad-Schlüssel, RefCount, Fan-out und Backpressure umsetzen.":
    "Zusammengesetzte ClientSubscriptionKeys, RefCount, geordnete Accepted/Rejected-Ereignisse, per-Client-Fan-out und echte begrenzte Backpressure umsetzen.",
    "ID-Validierung, Resolver, Lease-Grenzen, Korrelation, Subscription-RefCount.":
    "ID-Validierung, TreePathResolver, DeviceKey/DeviceId-Rückabbildung, atomare Lease-Grenzen, Korrelation, Subscription-RefCount und Ereignisreihenfolge.",
    "ValueChanged sowie SubscriptionCreated/Disposed gehören zum MVP. ChildNodeAdded/Removed sind laut Ausgangskonzept ebenfalls Pflicht und benötigen eine verbindliche Erweiterung der gemeinsamen Tree-/Container-Basis.":
    "ValueChanged, SubscriptionCreated/Disposed und ChildNodeAdded/Removed gehören laut Konzept zum MVP. Zusätzlich ist zu bestätigen, ob WatchNodeEvents die bisher getrennte Subscribe-Operation bewusst als langlebigen Aufruf bündelt.",
}

counts = {key: 0 for key in replacements}

def replace_paragraph(paragraph):
    text = paragraph.text
    if text not in replacements:
        return
    new_text = replacements[text]
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)
    counts[text] += 1

for paragraph in doc.paragraphs:
    replace_paragraph(paragraph)

for table in doc.tables:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if not tr_pr.xpath('./w:cantSplit'):
            tr_pr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_paragraph(paragraph)
    if table.rows:
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        if not header_pr.xpath('./w:tblHeader'):
            header_pr.append(OxmlElement('w:tblHeader'))

missing = [text for text, count in counts.items() if count != 1]
if missing:
    raise RuntimeError("Expected exactly one replacement for:\n" + "\n---\n".join(f"count={counts[text]}: {text}" for text in missing))

doc.save(path)
print(path)
print(f"replacements={sum(counts.values())}; tables={len(doc.tables)}")