from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Inches
from PIL import Image

base = Path(r"C:\zenlab-Projects\ba-thesis\implementation-concept")
src = base / "DeviceIntegrationPlugin_Kurzkonzept_geprueft.docx"
out = base / "DeviceIntegrationPlugin_Kurzkonzept_kompakt.docx"
staged = base / ".render" / "compact_staged.docx"
image = base / "diagrams" / "07_gesamtklassendiagramm_implementierung.png"

doc = Document(src)
replacements = {
    "Abbildung 3: Getrennte Service-/Proxy-Paare für Client- und Agent-Seite (PlantUML)":
    "Abbildung 3: Kompaktes Klassendiagramm des MVP-Kerns (PlantUML)",
    "DeviceIntegrationClientService und DeviceIntegrationAgentService laufen im Plugin. Sie greifen ausschließlich über Lifecycle-, Dispatcher-, Connection- und Subscription-Abstraktionen auf die gemeinsam als SingleInstance registrierten Zustände zu. Weil ein vorhandener DeviceAgent SubscriptionCreated bereits synchron vor seiner MethodResponse melden kann, puffert der SubscriptionManager diese Reihenfolge im Pending-Zustand begrenzt und gibt an Clients stets Accepted oder Rejected vor Created aus.":
    "DeviceIntegrationClientService und DeviceIntegrationAgentService bleiben dünne gRPC-Endpunkte. Der ClientService delegiert an Registry, Dispatcher und SubscriptionManager; der AgentService an Lifecycle und ConnectionManager. Fachlogik liegt damit in kleinen Use-Case-Klassen statt in Services oder einem allgemeinen Manager.",
    "Gerätespezifische Protokolle und TreeNode-Logik bleiben in den bestehenden DeviceAgents. Die gemeinsame Management-Basis beziehungsweise der DeviceAgent-Composition-Root muss jedoch vom alten DeviceManagementPlugin entkoppelt werden.":
    "Gerätespezifische Protokolle und TreeNode-Logik bleiben in den bestehenden DeviceAgents. Für den MVP gilt YAGNI: eine In-Memory-Registry, kein Gateway, kein neuer WebAPI-Adapter und keine zusätzlichen Manager-Schichten. Interfaces werden nur an echten Prozess- oder Modulgrenzen verwendet; interne Helper werden erst bei einer klaren eigenen Verantwortung extrahiert. Abhängigkeiten werden im Composition Root verdrahtet, nicht über Service-Locator oder globale Zustände.",
}
counts = {key: 0 for key in replacements}
for paragraph in doc.paragraphs:
    if paragraph.text in replacements:
        old = paragraph.text
        new = replacements[old]
        if paragraph.runs:
            paragraph.runs[0].text = new
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(new)
        counts[old] += 1

missing = [key for key, count in counts.items() if count != 1]
if missing:
    raise RuntimeError("Expected one replacement each: " + repr([(key, counts[key]) for key in missing]))

with Image.open(image) as im:
    ratio = im.width / im.height
width = 6.1
height = width / ratio
shape = doc.inline_shapes[2]
shape.width = Inches(width)
shape.height = Inches(height)

doc.save(staged)
with ZipFile(staged, "r") as zin, ZipFile(out, "w", ZIP_DEFLATED) as zout:
    for info in zin.infolist():
        data = image.read_bytes() if info.filename == "word/media/image3.png" else zin.read(info.filename)
        zout.writestr(info, data)

print(out)
print(f"replacements={sum(counts.values())}; figure3={width:.2f}x{height:.2f} in")