from pathlib import Path
from docx import Document
path = Path(r"C:\zenlab-Projects\ba-thesis\implementation-concept\DeviceIntegrationPlugin_Kurzkonzept_kompakt.docx")
doc = Document(path)
heading_count = 0
note_count = 0
for index, paragraph in enumerate(doc.paragraphs):
    if paragraph.text == "4.1 Empfohlene Projekte":
        paragraph.runs[0].text = "4.1 Logische Module"
        heading_count += 1
        for next_paragraph in doc.paragraphs[index + 1:]:
            if next_paragraph.text == "":
                next_paragraph.add_run("Die Modulnamen beschreiben Verantwortungs- und Abhängigkeitsgrenzen. Eine eigene Assembly ist erst nötig, wenn eine Grenze technisch durchgesetzt werden muss; für den MVP dürfen zusammengehörige Plugin-Klassen in einem Projekt bleiben.")
                note_count += 1
                break

cell_replacements = {
    "DeviceIntegration.Plugin.Core": "DeviceIntegration.Plugin",
    "Registry, Lifecycle, SessionTerminator, Connections, Dispatcher, Services und Subscription-Verwaltung.":
    "MEF-/DI-Hosting, gRPC-Services, Registry, Lifecycle, Connections, Dispatcher und Subscriptions.",
}
cell_counts = {key: 0 for key in cell_replacements}
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if paragraph.text in cell_replacements:
                    old = paragraph.text
                    paragraph.runs[0].text = cell_replacements[old]
                    for run in paragraph.runs[1:]:
                        run.text = ""
                    cell_counts[old] += 1

if heading_count != 1 or note_count != 1 or any(count != 1 for count in cell_counts.values()):
    raise RuntimeError((heading_count, note_count, cell_counts))
doc.save(path)
print(path)